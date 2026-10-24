"""Write real Office test documents, as an editor save would."""
import sys
from pathlib import Path
from docx import Document
from pptx import Presentation
from openpyxl import Workbook

root = Path(sys.argv[2]) if len(sys.argv) > 2 else Path(__file__).parent / "documents"
root.mkdir(exist_ok=True)
marker = sys.argv[1]
document = Document()
document.add_paragraph(f"The officewatch project status is {marker}. Word briefing.")
document.save(root / "report.docx")
slides = Presentation()
slide = slides.slides.add_slide(slides.slide_layouts[6])
slide.shapes.add_textbox(0, 0, 3000000, 1000000).text = f"The officewatch project status is {marker}. PowerPoint briefing."
slides.save(root / "report.pptx")
sheet = Workbook()
sheet.active.append([f"The officewatch project status is {marker}. Excel briefing."])
sheet.save(root / "report.xlsx")
(root / "control.txt").write_text(f"The officewatch project status is {marker}. Text control.")
print(f"Saved real DOCX/PPTX/XLSX/TXT documents with {marker}")
